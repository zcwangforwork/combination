"""
test_template_controlled.py - 受控形态模板层测试（方案模块四 §4.5/§5.2/§5.3）

覆盖:
- 封面受控状态行: 默认"受控" / env DOC_CONTROLLED_STATUS 三态覆盖
- 封面编制依据行: 法规名 + GB/T 42061-2022 常挂（64号/107号 双轨均含）
- 修订记录: 5 列表格、首行 V1.0"新增"、升版条目自动追加（修订人/生效日期留空）
- 正文节页眉三要素: 文件编号 | 标题缩略 | 版本+受控状态
- doc_version 参数透传（封面与修订记录首行一致）
"""
import sys
from pathlib import Path

import pytest
from docx import Document

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root))

SAMPLE_MARKDOWN = """# 设计输入文件

## 1 产品概述

输注精度应满足产品规格要求。
"""


@pytest.fixture(autouse=True)
def clean_env(monkeypatch):
    for k in ("DOC_CONTROLLED_STATUS", "DOC_VERSION", "DOC_MAKER",
              "DOC_REVIEWER", "DOC_APPROVER", "DOC_SIGN_DATE"):
        monkeypatch.delenv(k, raising=False)


def _build(doc_version="", revision_entries=None, doc_type="design_input",
           product_name="贴敷式胰岛素泵A7"):
    from app.services.template import TemplateService
    ts = TemplateService()
    doc = ts.load_template(doc_type)
    ts.fill_template(doc, SAMPLE_MARKDOWN, product_name, doc_type,
                     doc_version=doc_version, revision_entries=revision_entries)
    return doc


def _cover_info_table(doc):
    for t in doc.tables:
        if len(t.columns) == 4 and len(t.rows) >= 12:
            return t
    return None


def _revision_table(doc):
    for t in doc.tables:
        if len(t.columns) == 5:
            return t
    return None


def _body_header_text(doc) -> str:
    return doc.sections[-1].header.paragraphs[0].text


def test_cover_controlled_status_default():
    doc = _build()
    info = _cover_info_table(doc)
    assert info is not None
    assert info.cell(4, 0).text.startswith("受控状态")
    assert info.cell(4, 1).text.strip() == "受控"


def test_cover_controlled_status_env_override(monkeypatch):
    monkeypatch.setenv("DOC_CONTROLLED_STATUS", "非受控副本")
    doc = _build()
    info = _cover_info_table(doc)
    assert info.cell(4, 1).text.strip() == "非受控副本"
    # 正文节页眉同步受控状态三态
    assert "非受控副本" in _body_header_text(doc)


def test_cover_regulatory_basis_line():
    doc = _build()
    info = _cover_info_table(doc)
    assert info.cell(5, 0).text.startswith("编制依据")
    basis = info.cell(5, 1).text
    # 64号/107号 双轨均包含法规名 + GB/T 42061 常挂（断言与施行日无关的稳定部分）
    assert "医疗器械生产质量管理规范" in basis
    assert "GB/T 42061-2022" in basis


def test_revision_table_first_row_and_entries():
    from app.services.doc_control import build_revision_entry
    ent = build_revision_entry("V1.1", "用户指令", "修改流速参数要求")
    doc = _build(doc_version="V1.1", revision_entries=[ent])

    rt = _revision_table(doc)
    assert rt is not None
    # 首行：版本 + "新增"
    first = [c.text.strip() for c in rt.rows[1].cells]
    assert first[0] == "V1.1" and first[2] == "新增"
    # 升版条目行：版本/来源/概述齐备，修订人/生效日期留空待手填
    row = [c.text.strip() for c in rt.rows[2].cells]
    assert row[0] == "V1.1" and row[1] == "用户指令"
    assert row[2] == "修改流速参数要求"
    assert row[3] == "" and row[4] == ""


def test_body_header_three_elements():
    doc = _build()
    hdr = _body_header_text(doc)
    assert "V1.0｜受控" in hdr
    info = _cover_info_table(doc)
    doc_no = info.cell(1, 1).text.strip()
    assert doc_no and doc_no in hdr


def test_doc_version_param_propagates():
    doc = _build(doc_version="V2.3")
    info = _cover_info_table(doc)
    assert info.cell(2, 1).text.strip() == "V2.3"
    rt = _revision_table(doc)
    assert [c.text.strip() for c in rt.rows[1].cells][0] == "V2.3"
    assert "V2.3｜受控" in _body_header_text(doc)
