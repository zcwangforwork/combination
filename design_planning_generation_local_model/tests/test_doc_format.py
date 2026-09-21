"""文档格式工具测试：set_document_format / _apply_format_overrides

覆盖：
1. 工具：增量合并（只改传入项）/ 空参数不变 / 旁路写入
2. template._apply_format_overrides：正文字体/字号/行距/页边距实际写入 docx
"""
import asyncio
import json

from app.services import agent_tools


def test_set_document_format_merges():
    agent_tools.set_current_document_format({})
    agent_tools._pending_document_format.pop("format", None)

    async def _run():
        raw1 = await agent_tools.set_document_format.ainvoke({"body_font": "仿宋_GB2312"})
        raw2 = await agent_tools.set_document_format.ainvoke({
            "body_size_pt": 12, "line_spacing": 1.5, "margin_cm": 2.5,
        })
        return json.loads(raw1), json.loads(raw2)

    d1, d2 = asyncio.run(_run())
    assert d1["status"] == "ok" and "正文字体=仿宋_GB2312" in d1["updated"]
    # 增量合并：第二次设置保留第一次的字体
    assert d2["current_format"]["body_font"] == "仿宋_GB2312"
    assert d2["current_format"]["body_size_pt"] == 12
    assert d2["current_format"]["line_spacing"] == 1.5
    assert d2["current_format"]["margin_cm"] == 2.5
    # 旁路写入
    assert agent_tools._pending_document_format.get("format") == d2["current_format"]


def test_set_document_format_noop():
    agent_tools.set_current_document_format({"body_font": "宋体"})
    agent_tools._pending_document_format.pop("format", None)

    async def _run():
        raw = await agent_tools.set_document_format.ainvoke({})
        return json.loads(raw)

    d = asyncio.run(_run())
    assert d["status"] == "ok"
    assert d["current_format"] == {"body_font": "宋体"}
    assert agent_tools._pending_document_format.get("format") is None  # 无更新不写旁路


def test_apply_format_overrides_to_docx():
    from docx import Document
    from docx.shared import Cm
    from docx.oxml.ns import qn
    from app.services.template import TemplateService

    doc = Document()
    p = doc.add_paragraph("正文内容测试")
    doc.add_heading("标题测试", level=1)

    svc = TemplateService()
    svc._apply_format_overrides(doc, {
        "body_font": "仿宋_GB2312",
        "body_size_pt": 12,
        "heading_font": "楷体_GB2312",
        "line_spacing": 1.5,
        "margin_cm": 2.5,
    })

    # 正文字体（run 级 eastAsia）与字号
    body_run = p.runs[0]
    rfonts = body_run._element.rPr.rFonts
    assert rfonts.get(qn("w:eastAsia")) == "仿宋_GB2312"
    assert body_run.font.size.pt == 12

    # 行距
    assert p.paragraph_format.line_spacing == 1.5

    # 页边距
    sec = doc.sections[0]
    assert abs(sec.top_margin.cm - 2.5) < 0.01
    assert abs(sec.left_margin.cm - 2.5) < 0.01

    # 标题字体（样式级 eastAsia）
    h_style = doc.styles["Heading 1"]
    hrfonts = h_style.element.rPr.rFonts
    assert hrfonts.get(qn("w:eastAsia")) == "楷体_GB2312"


def test_apply_format_overrides_empty_noop():
    from docx import Document
    from app.services.template import TemplateService

    doc = Document()
    p = doc.add_paragraph("内容")
    before_size = p.runs[0].font.size  # None（未显式设置）
    TemplateService()._apply_format_overrides(doc, {})
    assert p.runs[0].font.size == before_size  # 空 overrides 不改动
